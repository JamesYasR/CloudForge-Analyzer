# -*- coding: utf-8 -*-
"""PPTX 组装工具箱（python-pptx）。"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image
import copy, os

SW, SH = 13.333, 7.5

INK = RGBColor(0x1B, 0x27, 0x33)
SUB = RGBColor(0x5A, 0x66, 0x72)
FAINT = RGBColor(0x98, 0xA4, 0xB0)
LINE = RGBColor(0xCB, 0xD5, 0xDE)
GRAY_L = RGBColor(0xF4, 0xF7, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

BLUE = RGBColor(0x1F, 0x5C, 0x99)
BLUE_L = RGBColor(0xE3, 0xED, 0xF7)
BLUE_D = RGBColor(0x16, 0x44, 0x70)

RED = RGBColor(0xC0, 0x39, 0x2B)
RED_L = RGBColor(0xFB, 0xE7, 0xE4)

# --- 兼容别名 ---
ORANGE, ORANGE_L = RED, RED_L
GREEN, GREEN_L = BLUE, BLUE_L
PURPLE, PURPLE_L = BLUE, BLUE_L
YELLOW_L, YELLOW_D = BLUE_L, BLUE

LATIN = "Microsoft YaHei"
EA = "微软雅黑"


def set_font(run, size=18, bold=False, color=INK, italic=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    f.name = LATIN
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {})
            rPr.append(e)
        e.set("typeface", EA)


def textbox(slide, x, y, w, h, lines, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT,
            line_spacing=1.3, anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if isinstance(lines, str):
        lines = [lines]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        if isinstance(ln, str):
            ln = {"t": ln}
        r = p.add_run()
        r.text = ln["t"]
        set_font(r, ln.get("size", size), ln.get("bold", bold), ln.get("color", color),
                 ln.get("italic", False))
        if "space_before" in ln:
            p.space_before = Pt(ln["space_before"])
    return tb


def rect(slide, x, y, w, h, fill=None, line=None, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
         radius=0.06, dash=None):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sh.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
        if dash:
            sh.line.dash_style = dash
    sh.shadow.inherit = False
    sh.text_frame.word_wrap = True
    return sh


def em_width_in(text, size):
    """估算文本宽度（英寸）：CJK 1.0 em，拉丁/数字约 0.55 em。"""
    em = 0.0
    for ch in text:
        em += 1.0 if ord(ch) > 0x2E80 else 0.55
    return em * size / 72.0


def title(slide, text, sub=None, y=0.34, size=28):
    rect(slide, 0.5, y + 0.06, 0.075, 0.52, fill=BLUE, shape=MSO_SHAPE.RECTANGLE)
    textbox(slide, 0.72, y, 11.6, 0.6, text, size=size, bold=True, color=INK)
    if sub:
        x = 0.72 + em_width_in(text, size) + 0.34
        textbox(slide, x, y + 0.15, 12.83 - x, 0.4, sub, size=15, color=SUB)
    rect(slide, 0.5, y + 0.86, 12.33, 0.018, fill=LINE, shape=MSO_SHAPE.RECTANGLE)


def footer(slide, page, note="凹塘/凹坑测量（需求 1.7）· 思路与原理"):
    textbox(slide, 0.5, 7.06, 8.0, 0.3, note, size=11, color=FAINT)
    textbox(slide, 12.2, 7.06, 0.63, 0.3, str(page), size=11, color=FAINT, align=PP_ALIGN.RIGHT)


def picture_fit(slide, path, x, y, w, h, border=True, shadow=False):
    """等比缩放并居中放入 (x,y,w,h) 框内。返回实际放置矩形。"""
    iw, ih = Image.open(path).size
    ar = iw / ih
    if w / h > ar:
        ph, pw = h, h * ar
    else:
        pw, ph = w, w / ar
    px, py = x + (w - pw) / 2, y + (h - ph) / 2
    if border:
        rect(slide, px - 0.02, py - 0.02, pw + 0.04, ph + 0.04, fill=None, line=LINE, lw=0.75,
             shape=MSO_SHAPE.RECTANGLE)
    slide.shapes.add_picture(path, Inches(px), Inches(py), Inches(pw), Inches(ph))
    return px, py, pw, ph


def lead(slide, text, y=1.20, x=0.5, w=12.33, h=1.02):
    """页首“讲人话”的总起句：浅底 + 左侧色条。"""
    rect(slide, x, y, w, h, fill=GRAY_L, line=None, radius=0.08)
    rect(slide, x, y + 0.06, 0.06, h - 0.12, fill=BLUE, shape=MSO_SHAPE.RECTANGLE)
    textbox(slide, x + 0.3, y + 0.12, w - 0.6, h - 0.2, text, size=15, color=INK,
            line_spacing=1.34)


def placeholder(slide, x, y, w, h, label, hint=None, size=16):
    """占位图框：虚线边框 + 提示文字。"""
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    sh = rect(slide, x, y, w, h, fill=GRAY_L, line=FAINT, lw=1.5, dash=MSO_LINE_DASH_STYLE.DASH)
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "【占位图】" + label
    set_font(r, size, True, SUB)
    if hint:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        r2.text = hint
        set_font(r2, size - 3, False, FAINT)
    return sh


def table(slide, x, y, w, h, data, col_w=None, header=True, size=13, header_size=13,
          row_h=None):
    rows, cols = len(data), len(data[0])
    gt = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = gt.table
    for j, cw in enumerate(col_w or [w / cols] * cols):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(data):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = Inches(0.07)
            cell.margin_right = Inches(0.05)
            cell.margin_top = Inches(0.02)
            cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if header and i == 0:
                cell.fill.fore_color.rgb = BLUE
            else:
                cell.fill.fore_color.rgb = WHITE if i % 2 else GRAY_L
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER if (j > 0 or header) else PP_ALIGN.LEFT
            r = p.add_run()
            if isinstance(val, tuple):
                val, col = val
            else:
                col = WHITE if (header and i == 0) else INK
            r.text = str(val)
            set_font(r, header_size if (header and i == 0) else size,
                     header and i == 0, col)
    if row_h:
        for i in range(rows):
            tbl.rows[i].height = Inches(row_h)
    return tbl


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def new_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])
